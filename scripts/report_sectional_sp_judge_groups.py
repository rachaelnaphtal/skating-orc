#!/usr/bin/env python3
"""
National-style analysis workbook for a named Singles/Pairs sectional roster.

Uses the same analysis / raw_data / Lookup layout as the National SP Judge
report, with Section, Group, and Status in the identity columns. Championships
columns are omitted. Default window is 24-25, 25-26, and 26-27 (all activity
plus a matching sectionals block).

Example::

    python scripts/report_sectional_sp_judge_groups.py
    python scripts/report_sectional_sp_judge_groups.py -o analysisTemp/sectional_sp_groups.xlsx
"""

from __future__ import annotations

import argparse
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

_REPO = Path(__file__).resolve().parents[1]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from activityAnalysis.international_listing_seasons import (  # noqa: E402
    season_codes_ending_at,
)
from activityAnalysis.load_activity_data import (  # noqa: E402
    DISC_SINGLES_PAIRS_ID,
    activity_database_is_postgresql,
    calendar_years_for_usfs_season_codes,
)
from analytics import JudgeAnalytics  # noqa: E402
from database import get_db_session  # noqa: E402
from officials_competition_types import COMPETITION_SCOPE_ALL  # noqa: E402
from scripts.national_judge_report_thresholds import (  # noqa: E402
    SINGLES_PAIRS_THRESHOLDS,
)
from scripts.national_sp_judge_analysis_xlsx import (  # noqa: E402
    write_national_sp_judge_analysis_xlsx,
)
from scripts.report_national_singles_pairs_judge_errors import (  # noqa: E402
    REPORT_DISCIPLINE_CONFIGS,
    _activity_tracker_columns,
    _competition_ids_for_seasons,
    _deviation_marking_scores_for_scope,
    _marking_score_for_labels,
    _metric_columns,
    _metrics_for_competitions,
    _official_candidate_judge_ids_for_roster,
    _official_identity_labels,
    _performance_block_season_span_label,
    _pool_judge_stats,
    _pool_season_stats,
    _segment_counts_for_judge_ids,
    _segment_entries_by_judge,
    _typed_competitions_for_seasons,
)
from scripts.report_singles_pairs_judge_activity import (  # noqa: E402
    _resolve_official_ids,
)

DEFAULT_END_SEASON = 2627
DEFAULT_N_SEASONS = 3
DEFAULT_OUTPUT = Path("analysisTemp/sectional_sp_judge_group_report.xlsx")

GROUP_1 = "Group 1"
GROUP_2 = "Group 2"
GROUP_3 = "Group 3"
GROUP_4 = "Group 4"
GROUP_5 = "Group 5"
GROUP_ORDER: tuple[str, ...] = (GROUP_1, GROUP_2, GROUP_3, GROUP_4, GROUP_5)

# Fills match the source list blocks (light green, mid-green, lavender, yellow, pink).
GROUP_FILLS: dict[str, str] = {
    GROUP_1: "C6EFCE",
    GROUP_2: "A9D08E",
    GROUP_3: "D9D2E9",
    GROUP_4: "FFE699",
    GROUP_5: "F8CBAD",
}


@dataclass(frozen=True)
class RosterJudge:
    name: str
    group: str
    section: str
    status: str
    lookup_name: str | None = None

    def directory_query_name(self) -> str:
        return self.lookup_name or self.name


# Source list order. Aliases are directory names when the sheet name differs.
ROSTER: tuple[RosterJudge, ...] = (
    RosterJudge("Katie Beriau Wheatley", GROUP_1, "Pacific Coast", "Return", "Katie Beriau"),
    RosterJudge("Rhea Sy-Benedict", GROUP_1, "Pacific Coast", "Return"),
    RosterJudge("Olivia Molina", GROUP_1, "Midwestern", "New"),
    RosterJudge("Lyra Katzman", GROUP_1, "Eastern", "New"),
    RosterJudge("Ariel Davydov", GROUP_1, "Midwestern", "New"),
    RosterJudge("Samir Mallya", GROUP_1, "Pacific Coast", "New"),
    RosterJudge("Harrison Choate", GROUP_1, "Eastern", "New"),
    RosterJudge("Danielle Anzalone", GROUP_1, "Eastern", "New"),
    RosterJudge("Sebastien Payannet", GROUP_1, "Pacific Coast", "New"),
    RosterJudge("Elizabeth Wright-Johnson", GROUP_2, "Eastern", "New"),
    RosterJudge("Yuu Ohno", GROUP_2, "Pacific Coast", "New"),
    RosterJudge("Hayley Pangle", GROUP_2, "Eastern", "New"),
    RosterJudge("Kristine Brickel", GROUP_2, "Midwestern", "New", "Kristy Brickel"),
    RosterJudge("Isabella Yung", GROUP_3, "Pacific Coast", "New"),
    RosterJudge("Selena Li", GROUP_3, "Eastern", "New"),
    RosterJudge("Angela Chen", GROUP_3, "Pacific Coast", "New"),
    RosterJudge("Lindsay Quon", GROUP_3, "Pacific Coast", "New"),
    RosterJudge("Kelvin Li", GROUP_3, "Pacific Coast", "New"),
    RosterJudge("Onia Gelecinskyj", GROUP_3, "Midwestern", "New"),
    RosterJudge("Kristofer Ogren", GROUP_4, "Midwestern", "New"),
    RosterJudge("Holly Tanner", GROUP_4, "Pacific Coast", "New"),
    RosterJudge("Ginger Whatley", GROUP_5, "Eastern", "New"),
    RosterJudge("Lori Cochran", GROUP_5, "Eastern", "Return"),
    RosterJudge("George Chow", GROUP_5, "Pacific Coast", "Return"),
    RosterJudge("Leslie Amacker", GROUP_5, "Midwestern", "New"),
    RosterJudge("Sam Gordon", GROUP_5, "Eastern", "Return"),
)

RANK_SPECS: tuple[tuple[str, str, bool], ...] = (
    ("competition_count", "Group rank: comps", True),
    ("junior_senior_segment_count", "Group rank: Jr/Sr segs", True),
    ("anomaly_rate_pct", "Group rank: anomaly %", False),
    ("element_marking_score", "Group rank: element", False),
    ("pcs_marking_score", "Group rank: PCS", False),
    ("sectionals_anomaly_rate_pct", "Group rank: sect anomaly %", False),
    ("sectionals_element_marking_score", "Group rank: sect element", False),
    ("sectionals_pcs_marking_score", "Group rank: sect PCS", False),
)


def add_within_group_ranks(
    df: pd.DataFrame,
    *,
    group_col: str = "Group",
    specs: tuple[tuple[str, str, bool], ...] = RANK_SPECS,
) -> pd.DataFrame:
    """Add min-ranks within each group. ``higher_is_better=True`` ranks 1 as largest."""
    out = df.copy()
    for value_col, rank_col, higher_is_better in specs:
        if value_col not in out.columns:
            continue
        out[rank_col] = out.groupby(group_col, sort=False)[value_col].rank(
            method="min",
            ascending=not higher_is_better,
            na_option="keep",
        )
    return out


def _mean(series: pd.Series) -> float | None:
    vals = pd.to_numeric(series, errors="coerce").dropna()
    if vals.empty:
        return None
    return float(vals.mean())


def _section_mix(section: pd.Series) -> str:
    counts = section.value_counts()
    return ", ".join(f"{int(n)} {name}" for name, n in counts.items())


def group_comparison_frame(df: pd.DataFrame) -> pd.DataFrame:
    """One summary row per group, in source-list order."""
    rows: list[dict[str, Any]] = []
    present = [g for g in GROUP_ORDER if g in set(df["Group"].tolist())]
    for group in present:
        sub = df.loc[df["Group"] == group]
        rows.append(
            {
                "Group": group,
                "Judges": int(len(sub)),
                "New": int((sub["Status"] == "New").sum()),
                "Return": int((sub["Status"] == "Return").sum()),
                "Sections": _section_mix(sub["Section"]),
                "Mean comps": _mean(
                    sub.get("competition_count", sub.get("Total comps", pd.Series(dtype=float)))
                ),
                "Mean Jr/Sr segments": _mean(
                    sub.get(
                        "junior_senior_segment_count",
                        sub.get("Jr/Sr Segments", pd.Series(dtype=float)),
                    )
                ),
                "Mean anomaly %": _mean(sub.get("anomaly_rate_pct", pd.Series(dtype=float))),
                "Mean element marking": _mean(
                    sub.get("element_marking_score", pd.Series(dtype=float))
                ),
                "Mean PCS marking": _mean(
                    sub.get("pcs_marking_score", pd.Series(dtype=float))
                ),
                "Mean rule errors": _mean(
                    sub.get("total_rule_errors", pd.Series(dtype=float))
                ),
                "Mean last sectionals": _mean(
                    sub.get("last_sectionals_in_role", pd.Series(dtype=float))
                ),
                "Mean sect anomaly %": _mean(
                    sub.get("sectionals_anomaly_rate_pct", pd.Series(dtype=float))
                ),
                "Mean sect element marking": _mean(
                    sub.get("sectionals_element_marking_score", pd.Series(dtype=float))
                ),
                "Mean sect PCS marking": _mean(
                    sub.get("sectionals_pcs_marking_score", pd.Series(dtype=float))
                ),
            }
        )
    return pd.DataFrame(rows)


def resolve_roster(roster: tuple[RosterJudge, ...] = ROSTER) -> pd.DataFrame:
    """Map the source list to directory officials, preserving list order."""
    lookups = [entry.directory_query_name() for entry in roster]
    resolved = _resolve_official_ids(lookups)
    by_requested = {
        str(row["requested_name"]).strip().lower(): row
        for _, row in resolved.iterrows()
    }
    rows: list[dict[str, Any]] = []
    for index, entry in enumerate(roster):
        key = entry.directory_query_name().strip().lower()
        hit = by_requested.get(key)
        if hit is None:
            raise SystemExit(f"Could not resolve {entry.name!r} ({entry.directory_query_name()!r})")
        rows.append(
            {
                "roster_order": index,
                "requested_name": entry.name,
                "Name": hit["Name"],
                "official_id": int(hit["official_id"]),
                "full_name": hit["Name"],
                "Group": entry.group,
                "Section": entry.section,
                "Status": entry.status,
            }
        )
    return pd.DataFrame(rows)


def _official_mbr_numbers(official_ids: list[int]) -> dict[int, str]:
    if not official_ids:
        return {}
    with get_db_session() as session:
        rows = session.execute(
            text(
                """
                SELECT id, mbr_number
                FROM officials_analysis.officials
                WHERE id = ANY(:ids)
                """
            ),
            {"ids": [int(x) for x in official_ids]},
        ).all()
    return {int(oid): (mbr or "").strip() for oid, mbr in rows}


def _quality_for_roster(
    analytics: JudgeAnalytics,
    officials: pd.DataFrame,
    *,
    seasons: list[str],
) -> pd.DataFrame:
    config = REPORT_DISCIPLINE_CONFIGS["singles_pairs"]
    seg_discipline_ids = list(config.segment_discipline_type_ids)
    official_ids = [int(x) for x in officials["official_id"].tolist()]
    official_identity_labels = _official_identity_labels(
        analytics, roster_official_ids=official_ids
    )
    all_comp_ids = _competition_ids_for_seasons(
        analytics,
        seasons,
        competition_scope=COMPETITION_SCOPE_ALL,
    )
    season_span = _performance_block_season_span_label(seasons)
    sectionals_comp_ids, sectionals_seasons = _typed_competitions_for_seasons(
        analytics,
        seasons=seasons,
        officials_competition_type_ids=config.sectionals_type_ids,
    )
    print(
        f"Quality window {', '.join(seasons)}: {len(all_comp_ids)} competitions (all activity)"
    )
    print(
        f"  Sectionals ({season_span}): {len(sectionals_comp_ids)} competitions"
    )

    metrics = _metrics_for_competitions(
        analytics,
        competition_ids=all_comp_ids,
        seg_discipline_ids=seg_discipline_ids,
        seasons=seasons,
        competition_scope=COMPETITION_SCOPE_ALL,
        label="all activity",
        discipline_label=config.roster_label,
        include_rule_errors=True,
    )
    sectionals_metrics = _metrics_for_competitions(
        analytics,
        competition_ids=sectionals_comp_ids,
        seg_discipline_ids=seg_discipline_ids,
        seasons=sectionals_seasons,
        competition_scope=config.sectionals_competition_scope,
        label="sectionals",
        discipline_label=config.roster_label,
        include_rule_errors=True,
    )
    segment_entries_by_judge = _segment_entries_by_judge(
        analytics,
        competition_ids=all_comp_ids,
        seg_discipline_ids=seg_discipline_ids,
    )
    sectionals_segment_entries_by_judge = _segment_entries_by_judge(
        analytics,
        competition_ids=sectionals_comp_ids,
        seg_discipline_ids=seg_discipline_ids,
    )
    candidate_judge_ids_by_official = _official_candidate_judge_ids_for_roster(
        analytics, officials
    )
    pcs_marking, element_marking = _deviation_marking_scores_for_scope(
        analytics,
        seasons,
        seg_discipline_ids=seg_discipline_ids,
        competition_scope=COMPETITION_SCOPE_ALL,
        scope_label="all activity",
    )
    if seasons:
        sectionals_pcs_marking, sectionals_element_marking = (
            _deviation_marking_scores_for_scope(
                analytics,
                seasons,
                seg_discipline_ids=seg_discipline_ids,
                competition_scope=config.sectionals_competition_scope,
                scope_label=f"sectionals ({season_span})",
            )
        )
    else:
        sectionals_pcs_marking, sectionals_element_marking = {}, {}

    rows: list[dict[str, Any]] = []
    for _, off in officials.iterrows():
        oid = int(off["official_id"])
        labels = official_identity_labels.get(oid, [])
        label = labels[0] if labels else ""
        has_label = bool(labels)
        candidate_jids = candidate_judge_ids_by_official.get(oid, set())
        row: dict[str, Any] = {
            "official_id": oid,
            "judge_identity_label": label or None,
            "quality_seasons": ", ".join(seasons),
            "sectionals_seasons_included": ", ".join(sectionals_seasons),
        }
        row.update(
            _metric_columns(
                prefix="",
                totals=metrics["totals"],
                by_season=metrics["by_season"],
                seasons=seasons,
                label=label,
                has_label=has_label,
                seg=_segment_counts_for_judge_ids(
                    segment_entries_by_judge, candidate_jids
                )
                if candidate_jids
                else {},
                pcs_marking=pcs_marking,
                element_marking=element_marking,
                include_marking_scores=True,
                include_rule_errors=True,
                pooled_stats=_pool_judge_stats(metrics["totals"], labels)
                if has_label
                else None,
                pooled_by_season=_pool_season_stats(
                    metrics["by_season"], labels, seasons
                )
                if has_label
                else None,
                pooled_seg=_segment_counts_for_judge_ids(
                    segment_entries_by_judge, candidate_jids
                )
                if candidate_jids
                else None,
                pcs_marking_score=_marking_score_for_labels(pcs_marking, labels),
                element_marking_score=_marking_score_for_labels(
                    element_marking, labels
                ),
            )
        )
        row.update(
            _metric_columns(
                prefix="sectionals_",
                totals=sectionals_metrics["totals"],
                by_season=sectionals_metrics["by_season"],
                seasons=sectionals_seasons,
                label=label,
                has_label=has_label,
                seg=_segment_counts_for_judge_ids(
                    sectionals_segment_entries_by_judge, candidate_jids
                )
                if candidate_jids
                else {},
                pcs_marking=sectionals_pcs_marking,
                element_marking=sectionals_element_marking,
                include_marking_scores=True,
                include_rule_errors=True,
                pooled_stats=_pool_judge_stats(sectionals_metrics["totals"], labels)
                if has_label
                else None,
                pooled_by_season=_pool_season_stats(
                    sectionals_metrics["by_season"], labels, sectionals_seasons
                )
                if has_label
                else None,
                pooled_seg=_segment_counts_for_judge_ids(
                    sectionals_segment_entries_by_judge, candidate_jids
                )
                if candidate_jids
                else None,
                pcs_marking_score=_marking_score_for_labels(
                    sectionals_pcs_marking, labels
                ),
                element_marking_score=_marking_score_for_labels(
                    sectionals_element_marking, labels
                ),
            )
        )
        rows.append(row)
    return pd.DataFrame(rows)


def build_report(
    *,
    end_season: int = DEFAULT_END_SEASON,
    n_seasons: int = DEFAULT_N_SEASONS,
) -> tuple[pd.DataFrame, pd.DataFrame, list[int]]:
    season_codes = season_codes_ending_at(int(end_season), int(n_seasons))
    seasons = [str(c) for c in sorted(season_codes, reverse=True)]
    officials = resolve_roster(ROSTER)
    official_ids = [int(x) for x in officials["official_id"].tolist()]
    mbr_by_id = _official_mbr_numbers(official_ids)
    officials["mbr_number"] = officials["official_id"].map(mbr_by_id)
    config = REPORT_DISCIPLINE_CONFIGS["singles_pairs"]
    discipline_by_official = {oid: DISC_SINGLES_PAIRS_ID for oid in official_ids}

    with get_db_session() as session:
        analytics = JudgeAnalytics(session)
        quality = _quality_for_roster(analytics, officials, seasons=seasons)
        activity_cols = _activity_tracker_columns(
            session,
            official_ids,
            config=config,
            discipline_by_official=discipline_by_official,
            availability_title="",
            season_year_codes=season_codes,
        )

    summary = officials.merge(quality, on="official_id", how="left")
    for oid, act in activity_cols.items():
        mask = summary["official_id"] == int(oid)
        for key, val in act.items():
            summary.loc[mask, key] = val
    summary["activity_competition_count"] = summary.get("competition_count")
    summary["activity_segment_count"] = summary.get("segment_count")
    summary["activity_junior_senior_segment_count"] = summary.get(
        "junior_senior_segment_count"
    )
    summary["directory_name"] = summary["requested_name"]
    summary["section"] = summary["Section"]
    summary["group"] = summary["Group"]
    summary["status"] = summary["Status"]
    summary["international_judge"] = False
    summary["discipline"] = "singles_pairs"
    summary["activity_scope"] = "all activity"
    summary["seasons_included"] = ", ".join(seasons)
    summary = add_within_group_ranks(summary)
    summary = summary.sort_values("roster_order", kind="mergesort").reset_index(drop=True)
    comparison = group_comparison_frame(summary)
    return summary, comparison, season_codes


def _append_group_comparison_sheet(path: Path, comparison: pd.DataFrame) -> None:
    display = comparison.copy()
    for col in display.columns:
        if col.startswith("Mean "):
            display[col] = display[col].map(
                lambda v: None
                if v is None or (isinstance(v, float) and math.isnan(v))
                else round(float(v), 3)
            )
    wb = load_workbook(path)
    if "Group comparison" in wb.sheetnames:
        del wb["Group comparison"]
    ws = wb.create_sheet("Group comparison")
    ws["A1"] = (
        "Group means for the same all-activity and sectionals columns as analysis. "
        "Rows stay in source-list order."
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 28
    headers = list(display.columns)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(3, col_idx, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row_idx, record in enumerate(display.itertuples(index=False), start=4):
        group = str(record[0])
        fill = PatternFill("solid", fgColor=GROUP_FILLS.get(group, "FFFFFF"))
        for col_idx, value in enumerate(record, start=1):
            cell = ws.cell(row_idx, col_idx, None if pd.isna(value) else value)
            cell.fill = fill
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 18 if col_idx == 5 else 14
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["E"].width = 36
    wb.save(path)


def write_excel(
    path: Path,
    summary: pd.DataFrame,
    comparison: pd.DataFrame,
    *,
    season_codes: list[int],
) -> None:
    seasons = [str(c) for c in sorted(season_codes)]
    span = _performance_block_season_span_label(seasons)
    cal_years = calendar_years_for_usfs_season_codes(season_codes)
    sectionals_activity_min_year = min(cal_years) if cal_years else 0
    write_national_sp_judge_analysis_xlsx(
        summary,
        path,
        include_rule_errors=True,
        include_championships=False,
        identity_layout="sectional_roster",
        preserve_row_order=True,
        group_fills=GROUP_FILLS,
        total_comps_header="Total Comps (3 years) in Role",
        thresholds=SINGLES_PAIRS_THRESHOLDS,
        recent_period_header=f"Last 3 seasons ({span}, all activity)",
        performance_block_header=f"Competition Performance ({span}, all activity)",
        activity_column_label="Competition Activity",
        performance_analysis_header="Performance Analysis",
        sectionals_block_header=f"Sectionals ({span})",
        sectionals_performance_header="Sectionals Performance Analysis",
        sectionals_activity_min_year=sectionals_activity_min_year,
    )
    _append_group_comparison_sheet(path, comparison)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--end-season",
        type=int,
        default=DEFAULT_END_SEASON,
        help=f"Last USFS season code in the window (default {DEFAULT_END_SEASON}).",
    )
    parser.add_argument(
        "--seasons",
        type=int,
        default=DEFAULT_N_SEASONS,
        help=f"Number of seasons ending at --end-season (default {DEFAULT_N_SEASONS}).",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Output .xlsx path (default {DEFAULT_OUTPUT}).",
    )
    args = parser.parse_args()

    if not activity_database_is_postgresql():
        print(
            "This report requires PostgreSQL (DATABASE_URL). Activity tables are not "
            "available on SQLite.",
            file=sys.stderr,
        )
        return 1

    summary, comparison, season_codes = build_report(
        end_season=args.end_season,
        n_seasons=args.seasons,
    )
    write_excel(
        args.output,
        summary,
        comparison,
        season_codes=season_codes,
    )
    print(
        f"Wrote {args.output} ({len(summary)} officials, {len(comparison)} groups, "
        f"seasons {season_codes})"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

