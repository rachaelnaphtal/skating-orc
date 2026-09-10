from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd

_REPO = Path(__file__).resolve().parents[1]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))

from scripts.national_sp_judge_analysis_xlsx import ANALYSIS_SHEET
from scripts.report_sectional_sp_judge_groups import (
    GROUP_1,
    GROUP_2,
    GROUP_3,
    GROUP_4,
    GROUP_5,
    GROUP_FILLS,
    ROSTER,
    add_within_group_ranks,
    group_comparison_frame,
)
from scripts.national_sp_judge_analysis_xlsx import write_national_sp_judge_analysis_xlsx
from openpyxl import load_workbook


def test_roster_has_five_source_groups_in_list_order():
    names = [entry.name for entry in ROSTER]
    assert len(names) == 26
    assert len(set(names)) == 26
    groups = [entry.group for entry in ROSTER]
    assert groups.count(GROUP_1) == 9
    assert groups.count(GROUP_2) == 4
    assert groups.count(GROUP_3) == 6
    assert groups.count(GROUP_4) == 2
    assert groups.count(GROUP_5) == 5
    assert groups[0] == GROUP_1
    assert groups[-1] == GROUP_5


def test_roster_aliases_for_directory_names():
    by_name = {entry.name: entry for entry in ROSTER}
    assert by_name["Katie Beriau Wheatley"].directory_query_name() == "Katie Beriau"
    assert by_name["Kristine Brickel"].directory_query_name() == "Kristy Brickel"
    assert by_name["Selena Li"].directory_query_name() == "Selena Li"


def test_within_group_ranks_higher_activity_and_lower_anomaly():
    df = pd.DataFrame(
        {
            "Group": [GROUP_1, GROUP_1, GROUP_1, GROUP_2, GROUP_2],
            "Total comps": [10, 4, 10, 8, 2],
            "anomaly_rate_pct": [0.2, 1.5, 0.8, 3.0, 0.1],
        }
    )
    ranked = add_within_group_ranks(
        df,
        specs=(
            ("Total comps", "Group rank: comps", True),
            ("anomaly_rate_pct", "Group rank: anomaly %", False),
        ),
    )
    g1 = ranked.loc[ranked["Group"] == GROUP_1]
    assert g1["Group rank: comps"].tolist() == [1.0, 3.0, 1.0]
    assert g1["Group rank: anomaly %"].tolist() == [1.0, 3.0, 2.0]
    g2 = ranked.loc[ranked["Group"] == GROUP_2]
    assert g2["Group rank: comps"].tolist() == [1.0, 2.0]
    assert g2["Group rank: anomaly %"].tolist() == [2.0, 1.0]


def test_group_comparison_keeps_source_order_and_means():
    df = pd.DataFrame(
        {
            "Group": [GROUP_2, GROUP_1, GROUP_1],
            "Status": ["New", "Return", "New"],
            "Section": ["Eastern", "Pacific Coast", "Midwestern"],
            "Total comps": [4, 10, 6],
            "Jr/Sr Segments": [2, 8, 4],
            "Jr/Sr Starts": [20, 80, 40],
            "anomaly_rate_pct": [1.0, 0.5, 1.5],
        }
    )
    comparison = group_comparison_frame(df)
    assert comparison["Group"].tolist() == [GROUP_1, GROUP_2]
    g1 = comparison.iloc[0]
    assert g1["Judges"] == 2
    assert g1["New"] == 1
    assert g1["Return"] == 1
    assert g1["Mean comps"] == 8.0
    assert g1["Mean anomaly %"] == 1.0
    assert "Pacific Coast" in g1["Sections"]


def test_sectional_roster_workbook_has_section_and_hides_champs(tmp_path: Path):
    df = pd.DataFrame(
        [
            {
                "directory_name": "Samir Mallya",
                "section": "Pacific Coast",
                "group": GROUP_1,
                "status": "New",
                "appointment_year": 2023,
                "last_sectionals_in_role": 2025,
                "total_comps_in_role_2yr": 21,
                "activity_competition_count": 21,
                "activity_segment_count": 40,
                "activity_junior_senior_segment_count": 37,
                "competition_count": 21,
                "segment_count": 40,
                "junior_senior_segment_count": 37,
                "total_rule_errors": 1,
                "anomaly_rate_pct": 0.19,
                "element_marking_score": 0.82,
                "pcs_marking_score": 0.84,
                "sectionals_competition_count": 1,
                "sectionals_segment_count": 4,
                "sectionals_junior_senior_segment_count": 4,
                "sectionals_total_rule_errors": 0,
                "sectionals_anomaly_rate_pct": 0.05,
                "sectionals_element_marking_score": 0.83,
                "sectionals_pcs_marking_score": 0.90,
            }
        ]
    )
    out = tmp_path / "sectional_layout.xlsx"
    write_national_sp_judge_analysis_xlsx(
        df,
        out,
        include_championships=False,
        identity_layout="sectional_roster",
        preserve_row_order=True,
        group_fills=GROUP_FILLS,
        total_comps_header="Total Comps (3 years) in Role",
        recent_period_header="Last 3 seasons (24-25 through 26-27, all activity)",
        sectionals_block_header="Sectionals (24-25 through 26-27)",
        sectionals_activity_min_year=2024,
    )
    ws = load_workbook(out)[ANALYSIS_SHEET]
    assert ws["A3"].value == "Name"
    assert ws["B3"].value == "Section"
    assert ws["C3"].value == "Group"
    assert ws["D3"].value == "Status"
    assert ws["B4"].value == "Pacific Coast"
    assert ws["C4"].value == GROUP_1
    assert ws["D4"].value == "New"
    assert ws["H3"].value == "Total Comps (3 years) in Role"
    assert ws["H2"].value == "Last 3 seasons (24-25 through 26-27, all activity)"
    assert ws["R2"].value == "Sectionals (24-25 through 26-27)"
    assert ws.column_dimensions["F"].hidden is True
    assert ws.column_dimensions["AA"].hidden is True
    assert ws.column_dimensions["BB"].hidden is True
    assert ws["B4"].fill.fgColor.rgb.endswith(GROUP_FILLS[GROUP_1])
